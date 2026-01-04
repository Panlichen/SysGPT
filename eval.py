#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SysGPT 评估脚本
该脚本用于评估 SysGPT 模型在系统性能优化方法论预测任务上的表现
与 few-shot 基线模型进行比较，并计算精确率、召回率和 F1 分数
"""

import sys
import re
from openai import OpenAI
from openpyxl import load_workbook

# 示例文件路径，用于 few-shot 学习
EXAMPLE_FILE = './dataset/example_10'
# OpenAI API 密钥，需替换为实际有效的密钥
API_KEY = 'your api key for OpenAI'
# 微调模型密钥，需替换为实际的微调模型 ID
MODEL_KEY = 'your fine-tuned model key'


def run_sysgpt(client: OpenAI, problem_observation: str, model_key: str, temp: float):
    """
    使用 SysGPT 模型生成系统性能优化方案
    
    Args:
        client: OpenAI 客户端实例
        problem_observation: 问题描述和观察结果
        model_key: 微调模型的 ID
        temp: 生成温度，控制输出的随机性
    
    Returns:
        str: 模型生成的优化方案文本
    """
    # 创建聊天完成请求
    completion = client.chat.completions.create(
        model=model_key,
        temperature=temp,
        messages=[
            {
                "role": "system",
                "content": "You are an expert in Computer Science, especially in Systems area, who explains things specifically and comprehensively. You know the following categories that are common methodologies to improve system performance in a single-line execution, excluding those benefited from parallelism and algorithmic optimizations: 1.  Batching: Merge duplicate costs by grouping data or operations; 2. Caching: Memorize computed result and reuse it to avoid redundant computation; 3. Precomputing: Conduct initialization or execution in advance; 4. Deferring: Delay initialization or execution until it is needed or it has better context to make decision; 5. Relaxation: Cut workload size by sacrificing accuracy with approximation; 6. Contextualization: Collect additional data at runtime to make better decisions; 7. Hardware: Utilize specific hardware features, e.g., NUMA, NVM, FPGA, SmartNIC, to optimize workload computation; 8. Bypass: Skip existing layer by taking a fast path; 9. Delayering: Merge multiple layers into one to avoid intermediate costs among layers; 10. Decoupling: Split one layer into multiple layers to have finer control. Given problem description with observations, provide a system solution for improving performance to the problem. Explain the solution in detail using the methodologies described above."
            },
            {
                "role": "user",
                "content": f"Given problem description with observations, provide **solutions** for improving performance to the problem. Give solutions in bullet points. Each bullet item must have following format:\n\n`1. [Methodology] Expanded description of this optimization technique.`\n\n The **Methodology** must be selected from one of the 10 methodologies.\nHere is the problem description with observations:\n```\n{problem_observation}\n```\n"
            }
        ]
    )
    
    # 返回模型生成的内容
    return completion.choices[0].message.content


def run_few_shot(client: OpenAI, problem_observation: str, examples: str, temp: float):
    """
    使用 few-shot 方法生成系统性能优化方案（基线模型）
    
    Args:
        client: OpenAI 客户端实例
        problem_observation: 问题描述和观察结果
        examples: 示例优化方案文本
        temp: 生成温度，控制输出的随机性
    
    Returns:
        str: 模型生成的优化方案文本
    """
    # 创建聊天完成请求，使用 gpt-4o-2024-08-06 模型作为基线
    completion = client.chat.completions.create(
        model="gpt-4o-2024-08-06",
        temperature=temp,
        messages=[
            {
                "role": "system",
                "content": "You are an expert in Computer Science, especially in Systems area, who explains things specifically and comprehensively. You know the following categories that are common methodologies to improve system performance in a single-line execution, excluding those benefited from parallelism and algorithmic optimizations: 1.  Batching: Merge duplicate costs by grouping data or operations; 2. Caching: Memorize computed result and reuse it to avoid redundant computation; 3. Precomputing: Conduct initialization or execution in advance; 4. Deferring: Delay initialization or execution until it is needed or it has better context to make decision; 5. Relaxation: Cut workload size by sacrificing accuracy with approximation; 6. Contextualization: Collect additional data at runtime to make better decisions; 7. Hardware: Utilize specific hardware features, e.g., NUMA, NVM, FPGA, SmartNIC, to optimize workload computation; 8. Bypass: Skip existing layer by taking a fast path; 9. Delayering: Merge multiple layers into one to avoid intermediate costs among layers; 10. Decoupling: Split one layer into multiple layers to have finer control. Given problem description with observations, provide a system solution for improving performance to the problem. Explain the solution in detail using the methodologies described above."
            },
            {
                "role": "user",
                "content": f"Given problem description with observations, provide **solutions** for improving performance to the problem. Give solutions in bullet points. Each bullet item must have following format:\n\n`1. [Methodology] Expanded description of this optimization technique.`\n\n The **Methodology** must be selected from one of the 10 methodologies.\n{examples}\nHere is the problem description with observations:\n```\n{problem_observation}\n```\n"
            }
        ]
    )
    
    # 返回模型生成的内容
    return completion.choices[0].message.content


def extract_and_convert_methodologies(text):
    """
    从文本中提取优化方法论标签，并转换为向量表示
    
    Args:
        text: 包含优化方案的文本，其中方法论用 [Methodology] 格式标记
    
    Returns:
        list 或 bool: 成功时返回方法论向量，失败时返回 False
    """
    # 允许的方法论列表
    allowed = [
        "Batching", "Caching", "Precomputing", "Deferring", "Relaxation",
        "Contextualization", "Hardware", "Bypass", "Delayering", "Decoupling"
    ]
    
    # 方法论归一化映射，处理大小写和变体
    normalization_map = {
        "bypassing": "Bypass",
        "bypass": "Bypass",
        "deferring": "Deferring",
        "decoupling": "Decoupling",
        "delayering": "Delayering",
        "caching": "Caching",
        "batching": "Batching",
        "precomputing": "Precomputing",
        "contextualization": "Contextualization",
        "relaxation": "Relaxation",
        "hardware": "Hardware"
    }
    
    # 使用正则表达式提取所有 [Methodology] 标签
    tags = re.findall(r'\[(.*?)\]', text)
    
    # 验证并归一化标签
    normalized_tags = []
    for tag in tags:
        key = tag.lower()
        if key in normalization_map:
            normalized_tags.append(normalization_map[key])
        else:
            # 如果遇到无效标签，返回 False
            return False
    
    # 初始化方法论向量，长度为允许的方法论数量
    final_vector = [0] * len(allowed)
    
    # 将归一化后的标签转换为向量（one-hot 编码）
    for tag in normalized_tags:
        index = allowed.index(tag)
        final_vector[index] = 1
    
    return final_vector


def calculate_metrics(ground_truth, prediction):
    """
    计算预测结果的评估指标
    
    Args:
        ground_truth: 真实标签向量
        prediction: 预测标签向量
    
    Returns:
        tuple: 包含精确率、召回率和 F1 分数的元组
    """
    # 确保真实标签和预测标签长度相同
    assert len(ground_truth) == len(prediction), "Lists must be the same length"
    
    # 计算真阳性、假阳性和假阴性
    true_positive = sum((g == 1 and p == 1) for g, p in zip(ground_truth, prediction))
    false_positive = sum((g == 0 and p == 1) for g, p in zip(ground_truth, prediction))
    false_negative = sum((g == 1 and p == 0) for g, p in zip(ground_truth, prediction))
    
    # 计算精确率：TP / (TP + FP)
    precision = true_positive / (true_positive + false_positive) if (true_positive + false_positive) else 0
    # 计算召回率：TP / (TP + FN)
    recall = true_positive / (true_positive + false_negative) if (true_positive + false_negative) else 0
    # 计算 F1 分数：2 * (精确率 * 召回率) / (精确率 + 召回率)
    f1_score = (2 * precision * recall) / (precision + recall) if (precision + recall) else 0
    
    return precision, recall, f1_score


def run_test(client: OpenAI, temp: float, trial: int):
    """
    运行评估测试，比较 SysGPT 和 few-shot 基线模型的性能
    
    Args:
        client: OpenAI 客户端实例
        temp: 生成温度，控制输出的随机性
        trial: 每个问题的尝试次数，取最优结果
    """
    # 读取示例文件内容
    example_file = open(EXAMPLE_FILE, 'r').read()
    
    # 加载 Excel 数据集
    wb = load_workbook(filename='./dataset/dataset.xlsx')
    ws = wb['testset']  # 选择测试集工作表
    
    # 方法论列表，与 extract_and_convert_methodologies 函数中的 allowed 一致
    methodology = [
        "Batching", "Caching", "Precomputing", "Deferring", "Relaxation",
        "Contextualization", "Hardware", "Bypass", "Delayering", "Decoupling"
    ]
    
    # 初始化计数器和累加器
    n = 0  # 有效测试样本数量
    
    # SysGPT 模型的总指标
    sysgpt_total_precision = 0
    sysgpt_total_recall = 0
    sysgpt_total_f1 = 0
    
    # few-shot 基线模型的总指标
    few_shot_total_precision = 0
    few_shot_total_recall = 0
    few_shot_total_f1 = 0
    few_shot_method_num = 0  # few-shot 方法平均使用的方法论数量
    
    # 遍历测试集中的每一行（从第2行到第97行，共96个样本）
    for row in range(2, 98):
        # 检查该样本是否在测试范围内
        is_scope = ws['G' + str(row)].value
        
        # 如果标记为 'x'，跳过该样本
        if is_scope == 'x':
            continue
        # 如果标记为 'o'，处理该样本
        elif is_scope == 'o':
            print(f'=== {row} ===')
            
            # 获取问题描述和观察结果
            problem = ws['D' + str(row)].value
            observation = ws['E' + str(row)].value
            problem_observation = problem + '\n' + observation
            
            # 获取真实标签向量
            ground_truth = extract_and_convert_methodologies(ws['F' + str(row)].value)
            print(ground_truth)
            
            # ====== SysGPT 模型评估 ======
            
            # 初始化最佳指标
            best_f1 = -1
            best_precision = 0
            best_recall = 0
            
            # 多次尝试，取最优结果
            for i in range(0, trial):
                # 循环直到成功获取有效预测向量
                while True:
                    # 使用 SysGPT 模型生成答案
                    sysgpt_answer = run_sysgpt(client, problem_observation, MODEL_KEY, temp)
                    # 提取并转换方法论标签
                    sysgpt_prediction = extract_and_convert_methodologies(sysgpt_answer)
                    # 检查提取是否成功
                    if isinstance(sysgpt_prediction, bool):
                        print(sysgpt_answer)
                        print("Extract Error!! Retry...")
                    else:
                        break
                
                print("> Try ", i + 1)
                print(sysgpt_prediction)
                
                # 计算当前尝试的指标
                precision, recall, f1 = calculate_metrics(ground_truth, sysgpt_prediction)
                # 更新最佳指标
                if f1 > best_f1:
                    best_f1 = f1
                    best_precision = precision
                    best_recall = recall
            
            # 更新 SysGPT 模型的总指标
            sysgpt_total_precision += best_precision
            sysgpt_total_recall += best_recall
            sysgpt_total_f1 += best_f1
            n += 1
            
            # ====== 基线模型评估 ======
            
            # 初始化最佳指标
            best_f1 = -1
            best_precision = 0
            best_recall = 0
            best_num = 0
            
            # 多次尝试，取最优结果
            for i in range(0, trial):
                # 循环直到成功获取有效预测向量
                while True:
                    # 使用 few-shot 基线模型生成答案
                    few_shot = run_few_shot(client, problem_observation, example_file, temp)
                    # 提取并转换方法论标签
                    few_shot_prediction = extract_and_convert_methodologies(few_shot)
                    # 检查提取是否成功
                    if isinstance(few_shot_prediction, bool):
                        print(few_shot)
                        print("Extract Error!! Retry...")
                    else:
                        break
                
                print("> Try ", i + 1)
                print(few_shot_prediction)
                
                # 计算当前尝试的指标
                precision, recall, f1 = calculate_metrics(ground_truth, few_shot_prediction)
                # 更新最佳指标
                if f1 > best_f1:
                    best_f1 = f1
                    best_precision = precision
                    best_recall = recall
                    best_num = sum(few_shot_prediction)  # 计算使用的方法论数量
            
            # 更新 few-shot 基线模型的总指标
            few_shot_total_precision += best_precision
            few_shot_total_recall += best_recall
            few_shot_total_f1 += best_f1
            few_shot_method_num += best_num
            
            # 打印当前累计平均指标
            print("sysgpt  f1: ", sysgpt_total_f1 / n)
            print("fewshot f1: ", few_shot_total_f1 / n)
            print("fewshot # : ", few_shot_method_num / n)
    
    # 打印最终的平均指标
    print("=== sysgpt ===")
    print("avg_precision", sysgpt_total_precision / n)
    print("avg_recall", sysgpt_total_recall / n)
    print("avg_f1", sysgpt_total_f1 / n)
    
    print("=== few-shot ===")
    print("avg_precision", few_shot_total_precision / n)
    print("avg_recall", few_shot_total_recall / n)
    print("avg_f1", few_shot_total_f1 / n)


if __name__ == '__main__':
    """
    主函数，程序入口
    """
    # 初始化 OpenAI 客户端
    client = OpenAI(api_key=API_KEY)
    
    # 从命令行参数获取温度和尝试次数
    TEMP = float(sys.argv[1])
    TRIAL = int(sys.argv[2])
    
    # 打印测试信息
    print("Test info:")
    print(" - TEMP     : ", TEMP)
    print(" - Baseline : ", EXAMPLE_FILE)
    print(" - # trials : ", TRIAL)
    print("\n\n")
    
    # 运行测试
    run_test(client, TEMP, TRIAL)
